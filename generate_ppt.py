from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = Presentation()

def set_slide_bg(slide, rgb_color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = rgb_color

DARK_BLUE = RGBColor(15, 23, 42)
LIGHT_BLUE = RGBColor(241, 245, 249)
TEXT_DARK = RGBColor(15, 23, 42)
TEXT_LIGHT = RGBColor(241, 245, 249)
ACCENT = RGBColor(37, 99, 235)

def style_title(shape, text, is_dark_bg=False):
    shape.text = text
    p = shape.text_frame.paragraphs[0]
    p.font.color.rgb = TEXT_LIGHT if is_dark_bg else TEXT_DARK
    p.font.bold = True
    p.font.name = "Arial"

def add_bullet(text_frame, text, level=0, is_dark_bg=False):
    p = text_frame.add_paragraph()
    p.text = text
    p.level = level
    p.font.color.rgb = RGBColor(203, 213, 225) if is_dark_bg else RGBColor(71, 85, 105)
    p.font.name = "Arial"
    p.font.size = Pt(20 - (level * 2))

# 1. Title Slide
slide1 = prs.slides.add_slide(prs.slide_layouts[0])
set_slide_bg(slide1, DARK_BLUE)
style_title(slide1.shapes.title, "InvestorOS", True)
slide1.shapes.title.text_frame.paragraphs[0].font.size = Pt(60)

subtitle1 = slide1.placeholders[1]
subtitle1.text = "The Ultimate Investor & Growth Office Platform\nAutomate investor relations, board reporting, and growth metrics."
subtitle1.text_frame.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)
subtitle1.text_frame.paragraphs[0].font.size = Pt(24)

# 2. Executive Summary
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
set_slide_bg(slide2, LIGHT_BLUE)
style_title(slide2.shapes.title, "Executive Summary")
tf2 = slide2.placeholders[1].text_frame
tf2.text = "A Paradigm Shift in Founder-Investor Relations"
tf2.paragraphs[0].font.size = Pt(24)
tf2.paragraphs[0].font.color.rgb = ACCENT
add_bullet(tf2, "The modern startup operates at breakneck speed, but reporting is stuck in the past.")
add_bullet(tf2, "Founders and CFOs lose hundreds of hours annually compiling data manually.")
add_bullet(tf2, "InvestorOS is the operating system designed to bridge the gap between company performance and stakeholder communication.")
add_bullet(tf2, "Our goal: Enable founders to focus on building while we handle the reporting.")

# 3. The Problem Landscape
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
set_slide_bg(slide3, LIGHT_BLUE)
style_title(slide3.shapes.title, "The Problem Landscape")
tf3 = slide3.placeholders[1].text_frame
tf3.text = "Why is reporting so painful?"
tf3.paragraphs[0].font.size = Pt(24)
tf3.paragraphs[0].font.color.rgb = ACCENT
add_bullet(tf3, "Fragmented Data Sources", 0)
add_bullet(tf3, "Financials live in QuickBooks/Xero; growth metrics in Stripe/ChartMogul; relations in CRM.", 1)
add_bullet(tf3, "Manual Aggregation", 0)
add_bullet(tf3, "Teams spend days copying and pasting data into static spreadsheets and slides.", 1)
add_bullet(tf3, "Inconsistent Stakeholder Updates", 0)
add_bullet(tf3, "Investors receive irregular, unstandardized updates, leading to a lack of trust and alignment.", 1)

# 4. The Cost of Inaction
slide4 = prs.slides.add_slide(prs.slide_layouts[1])
set_slide_bg(slide4, DARK_BLUE)
style_title(slide4.shapes.title, "The Cost of Inaction", True)
tf4 = slide4.placeholders[1].text_frame
tf4.text = "What happens when reporting is broken?"
tf4.paragraphs[0].font.size = Pt(24)
tf4.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)
add_bullet(tf4, "Wasted Founder Time: ~20% of a founder's time is spent on administrative reporting tasks.", 0, True)
add_bullet(tf4, "Delayed Capital: Unprepared data rooms cause weeks of delay during crucial fundraising rounds.", 0, True)
add_bullet(tf4, "Poor Board Meetings: Meetings become status updates instead of strategic planning sessions.", 0, True)

# 5. Introducing InvestorOS
slide5 = prs.slides.add_slide(prs.slide_layouts[1])
set_slide_bg(slide5, LIGHT_BLUE)
style_title(slide5.shapes.title, "Introducing InvestorOS")
tf5 = slide5.placeholders[1].text_frame
tf5.text = "The Single Source of Truth"
tf5.paragraphs[0].font.size = Pt(24)
tf5.paragraphs[0].font.color.rgb = ACCENT
add_bullet(tf5, "A unified platform integrating directly with your core operational tools.")
add_bullet(tf5, "Real-time synchronization ensures your metrics are never out of date.")
add_bullet(tf5, "Designed specifically for scaling startups from Seed to Series C.")

# 6. Deep Dive: Dashboard & Growth Metrics
slide6 = prs.slides.add_slide(prs.slide_layouts[1])
set_slide_bg(slide6, LIGHT_BLUE)
style_title(slide6.shapes.title, "Dashboard & Growth Metrics")
tf6 = slide6.placeholders[1].text_frame
tf6.text = "Know your numbers instantly."
tf6.paragraphs[0].font.size = Pt(24)
tf6.paragraphs[0].font.color.rgb = ACCENT
add_bullet(tf6, "Live visibility into critical KPIs: ARR, MRR, Churn, CAC, LTV, and Cash Runway.")
add_bullet(tf6, "Customizable views for different stakeholders (Founders vs. Department Heads).")
add_bullet(tf6, "Automated cohort analysis and revenue retention tracking.")

# 7. Deep Dive: Investor Reporting
slide7 = prs.slides.add_slide(prs.slide_layouts[1])
set_slide_bg(slide7, LIGHT_BLUE)
style_title(slide7.shapes.title, "Automated Investor Reporting")
tf7 = slide7.placeholders[1].text_frame
tf7.text = "Professional updates in minutes."
tf7.paragraphs[0].font.size = Pt(24)
tf7.paragraphs[0].font.color.rgb = ACCENT
add_bullet(tf7, "Pre-built, best-practice templates for monthly and quarterly updates.")
add_bullet(tf7, "One-click data imports from your dashboard directly into the report.")
add_bullet(tf7, "Track investor engagement: see who opened your updates and which sections they read.")

# 8. Deep Dive: Board Deck Builder
slide8 = prs.slides.add_slide(prs.slide_layouts[1])
set_slide_bg(slide8, LIGHT_BLUE)
style_title(slide8.shapes.title, "Board Deck Builder")
tf8 = slide8.placeholders[1].text_frame
tf8.text = "Elevate your board meetings."
tf8.paragraphs[0].font.size = Pt(24)
tf8.paragraphs[0].font.color.rgb = ACCENT
add_bullet(tf8, "Say goodbye to manual copy-pasting into PowerPoint or Keynote.")
add_bullet(tf8, "Dynamically link charts and financial tables to your live data.")
add_bullet(tf8, "Standardized templates ensure you cover strategy, financials, and hiring seamlessly.")

# 9. Deep Dive: Fundraising Readiness
slide9 = prs.slides.add_slide(prs.slide_layouts[1])
set_slide_bg(slide9, LIGHT_BLUE)
style_title(slide9.shapes.title, "Fundraising Readiness & Scenario Modeling")
tf9 = slide9.placeholders[1].text_frame
tf9.text = "Be prepared for your next round, always."
tf9.paragraphs[0].font.size = Pt(24)
tf9.paragraphs[0].font.color.rgb = ACCENT
add_bullet(tf9, "Always-on Data Room: Keep your cap table, financials, and corporate docs organized.")
add_bullet(tf9, "Scenario Modeling: Stress-test your cash runway under different growth and hiring assumptions.")
add_bullet(tf9, "Share access securely with prospective investors with granular permissions.")

# 10. Value Proposition & ROI
slide10 = prs.slides.add_slide(prs.slide_layouts[1])
set_slide_bg(slide10, DARK_BLUE)
style_title(slide10.shapes.title, "Value Proposition & ROI", True)
tf10 = slide10.placeholders[1].text_frame
tf10.text = "Why companies choose InvestorOS"
tf10.paragraphs[0].font.size = Pt(24)
tf10.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)
add_bullet(tf10, "Time Savings: Recapture ~70% of time spent on administrative reporting.", 0, True)
add_bullet(tf10, "Increased Valuation: Clean data and professional reporting signal maturity to investors.", 0, True)
add_bullet(tf10, "Strategic Focus: Shift board meetings from data validation to strategic decision-making.", 0, True)

# 11. Conclusion / Q&A
slide11 = prs.slides.add_slide(prs.slide_layouts[0])
set_slide_bg(slide11, DARK_BLUE)
style_title(slide11.shapes.title, "Thank You", True)
slide11.shapes.title.text_frame.paragraphs[0].font.size = Pt(60)
subtitle11 = slide11.placeholders[1]
subtitle11.text = "Questions?\n\nClosing Statement: \"Don't let manual reporting slow down your growth. Let InvestorOS handle the data, so you can get back to building the future.\"\n\nLet's transform your investor relations."
subtitle11.text_frame.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)
subtitle11.text_frame.paragraphs[0].font.size = Pt(24)

prs.save("InvestorOS_Detailed_Presentation_v2.pptx")
print("Presentation successfully saved as InvestorOS_Detailed_Presentation_v2.pptx")
